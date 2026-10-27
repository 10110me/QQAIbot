import io
import zipfile


def parse_document(content: bytes, filename: str) -> list[dict]:
    """Run inside the indexer process; retain page numbers and reject archives/macros."""
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension in {"txt", "md", "csv", "json"}:
        if b"\x00" in content[:4096]:
            raise ValueError("Binary data is not a text document")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("gb18030")
        return [{"page": 1, "text": text[:1_000_000]}]
    if extension == "pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("PDF signature mismatch")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 500:
            raise ValueError("Encrypted or too many PDF pages")
        return [{"page": index + 1, "text": (page.extract_text() or "")[:100_000]} for index, page in enumerate(reader.pages)]
    if extension == "docx":
        if not content.startswith(b"PK"):
            raise ValueError("DOCX signature mismatch")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            info = archive.infolist()
            if len(info) > 2000 or sum(row.file_size for row in info) > 50_000_000:
                raise ValueError("DOCX decompression limit exceeded")
            if any("vbaProject" in row.filename or row.filename.startswith("/") or ".." in row.filename.split("/") for row in info):
                raise ValueError("Unsafe DOCX container")
        from docx import Document
        document = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows)
        return [{"page": 1, "text": text[:1_000_000]}]
    raise ValueError("Unsupported document type; executable, archive and macro formats denied")


def chunks(pages: list[dict], size: int = 900, overlap: int = 120) -> list[dict]:
    if size <= overlap or overlap < 0:
        raise ValueError("Invalid chunk parameters")
    result = []
    for page in pages:
        text = page["text"]
        for start in range(0, len(text), size-overlap):
            value = text[start:start+size].strip()
            if value:
                result.append({"page": page["page"], "text": value, "start": start})
    return result[:5000]
